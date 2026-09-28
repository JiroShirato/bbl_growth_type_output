from openpyxl.worksheet.worksheet import Worksheet

import one_pattern_class as opc
import init_values as iv
import ext_functions as ef


def calc_growth_patterns(in_gro_ex_list: list[str], in_att_list: list[int]) -> dict[str, opc.OnePattern]:
    """1つの成長期について、シートの条件に応じた全ての経験値パターンを計算

    Args:
        in_gro_ex_list (list[str]): CSVの1行（成長期名と各練習の最小値・最大値）
        in_att_list (list[int]): シートの条件（小AP, YUR, 集中・イマイチ, ギプスの有無, 出力対象）

    Returns:
        dict[str, opc.OnePattern]: パターン名をkeyとした全てのパターンの辞書

    Raises:
        ValueError: 数値に変換できない値がある、または大練習マイナスの値が空の場合
    """
    mini_ap, yur, concentrate, cast, output_type = in_att_list
    growth_name: str = in_gro_ex_list[0]

    patterns_dict: dict[str, opc.OnePattern] = create_base_patterns(in_gro_ex_list)

    # 1. 小APによる経験値追加処理
    if mini_ap == 1:
        add_value_to_train_patterns(patterns_dict, iv.MINI_AP_ADD_VALUE)

    # 2. APによる経験値倍増処理
    add_ap_patterns(patterns_dict)

    # 3. YURによる経験値追加処理
    if yur == 1:
        add_value_to_train_patterns(patterns_dict, iv.YUR_ADD_VALUE)

    # 4. メンタリスト能力持ちの嫁追加処理
    add_mentalist_patterns(patterns_dict)

    # 5. 集中、及びイマイチやる気が出ない…による経験値加減処理
    if concentrate != 0:
        apply_concentrate(patterns_dict, concentrate)

    # 6. 各鍛錬による補正
    add_training_patterns(patterns_dict, growth_name)

    # 7. 筋肉養成ギプスの適用
    if cast == 1:
        apply_cast(patterns_dict)

    # 8. 期待値算出（出力対象が期待値の時）
    if output_type == 1:
        for one_pattern in patterns_dict.values():
            one_pattern.calc_expected_value()

    return patterns_dict


def create_base_patterns(in_gro_ex_list: list[str]) -> dict[str, opc.OnePattern]:
    """CSVの1行から、補正前の基本パターン（小練習、大練習、大練習マイナス、自主トレ参加）を作成

    Args:
        in_gro_ex_list (list[str]): CSVの1行（成長期名と各練習の最小値・最大値）

    Returns:
        dict[str, opc.OnePattern]: パターン名をkeyとした基本パターンの辞書
    """
    # 基本パターンを格納する辞書
    patterns_dict: dict[str, opc.OnePattern] = {}

    # 小練習の値の取得
    small_opc: opc.OnePattern = opc.OnePattern.from_strings(in_gro_ex_list[1], in_gro_ex_list[2])
    patterns_dict["small_train_ex"] = small_opc

    # 大練習の値の取得
    large_opc: opc.OnePattern = opc.OnePattern.from_strings(in_gro_ex_list[3], in_gro_ex_list[4])
    patterns_dict["large_train_ex"] = large_opc

    # 大練習のマイナス値の取得
    minus_opc: opc.OnePattern = opc.OnePattern.from_strings(in_gro_ex_list[6], in_gro_ex_list[5])
    patterns_dict["large_minus_ex"] = minus_opc

    # 自主トレ参加の値の取得(小練習の値を使用、小練習の値がない場合はインスタンスだけ作る)
    small_opc_list: list[int] = small_opc.get_values_list()
    if len(small_opc_list) == 0:
        patterns_dict["participate_independent_training_ex"] = opc.OnePattern()
    else:
        independent_opc: opc.OnePattern = opc.OnePattern(min(small_opc_list) * 3, max(small_opc_list) * 3)
        patterns_dict["participate_independent_training_ex"] = independent_opc

    return patterns_dict


def add_value_to_train_patterns(in_ptn_dict: dict[str, opc.OnePattern], in_add_val: int) -> None:
    """大練習と小練習すべてのパターンに固定値を加算（小AP、YUR向け）

    Args:
        in_ptn_dict (dict[str, opc.OnePattern]): 1つの成長期のパターンの辞書
        in_add_val (int): 加算する値
    """
    for one_pattern_name, one_pattern in in_ptn_dict.items():
        if ef.is_train_pattern(one_pattern_name):
            one_pattern.add_ex_point(in_add_val)


def add_ap_patterns(in_ptn_dict: dict[str, opc.OnePattern]) -> None:
    """APによる経験値倍増のパターンを追加(倍率設定から、計算)

    Args:
        in_ptn_dict (dict[str, opc.OnePattern]): 1つの成長期のパターンの辞書
    """
    # 新たに追加する経験値パターンを格納する辞書
    add_ex_pattern_dict: dict[str, opc.OnePattern] = {}

    for ap_name, mul_val in iv.AP_MULTIPLICATION_FACTOR_DICT.items():
        for one_pattern_name, one_pattern in in_ptn_dict.items():
            # 大練習と小練習すべてが対象
            if ef.is_train_pattern(one_pattern_name):
                # APの名前の設定
                add_one_pattern_name: str = ap_name + "_" + one_pattern_name
                # 新しいパターンを作成し、APに設定された乗算値をかける（端数は切り捨て）
                add_ex_pattern_dict[add_one_pattern_name] = one_pattern.copy()
                add_ex_pattern_dict[add_one_pattern_name].mul_ex_point(mul_val)

    in_ptn_dict.update(add_ex_pattern_dict)


def add_mentalist_patterns(in_ptn_dict: dict[str, opc.OnePattern]) -> None:
    """メンタリスト能力持ちの嫁のパターンを追加

    Args:
        in_ptn_dict (dict[str, opc.OnePattern]): 1つの成長期のパターンの辞書
    """
    # 新たに追加する経験値パターンを格納する辞書
    add_ex_pattern_dict: dict[str, opc.OnePattern] = {}

    for one_pattern_name, one_pattern in in_ptn_dict.items():
        if one_pattern_name == "large_train_ex" or one_pattern_name == "small_train_ex":
            # 大練習と小練習が対象なら、メンタリストの名前"mental_"を追加
            add_one_pattern_name: str = "mental_" + one_pattern_name
        elif one_pattern_name.startswith("other_ap_"):
            # 1.5倍APなら、名前をメンタリスト"mental_"に変更
            add_one_pattern_name = one_pattern_name.replace("other_", "mental_")
        else:
            # 上記のパターン以外は対応しない
            continue
        # 新しいパターンを作成し、乗算値 iv.MENTALIST_WIFE_MUL_FACTOR をかける
        add_ex_pattern_dict[add_one_pattern_name] = one_pattern.copy()
        add_ex_pattern_dict[add_one_pattern_name].mul_ex_point(
            iv.MENTALIST_WIFE_MUL_FACTOR)

    in_ptn_dict.update(add_ex_pattern_dict)


def apply_concentrate(in_ptn_dict: dict[str, opc.OnePattern], in_conc_type: int) -> None:
    """集中、及びイマイチやる気が出ない…による経験値加減処理

    Args:
        in_ptn_dict (dict[str, opc.OnePattern]): 1つの成長期のパターンの辞書
        in_conc_type (int): 0: なし、1: 集中、2: イマイチ
    """
    for one_pattern_name, one_pattern in in_ptn_dict.items():
        # 大練習と小練習すべてが対象
        if ef.is_train_pattern(one_pattern_name):
            if in_conc_type == 1:
                # 集中の時
                one_pattern.add_concentrate_ex_point()
            elif in_conc_type == 2:
                # イマイチの時（2で割る）
                one_pattern.div_ex_point_with_ceil_and_floor(2)


def has_large_minus(in_ptn_dict: dict[str, opc.OnePattern], in_gro_name: str) -> bool:
    """大練習で経験値が下がるか（積極鍛錬と慎重鍛錬向け）

    Args:
        in_ptn_dict (dict[str, opc.OnePattern]): 1つの成長期のパターンの辞書
        in_gro_name (str): 成長期名（エラーメッセージ用）

    Returns:
        bool: 大練習マイナスがあれば True

    Raises:
        ValueError: 大練習マイナスの値が空の場合
    """
    # 大練習マイナスのパターンの値を格納したリスト
    temp_value_list: list[int] = in_ptn_dict["large_minus_ex"].get_values_list()
    try:
        return max(temp_value_list) < 0
    except ValueError as e:
        raise ValueError(
            f"大練習マイナスの値が不正です: {e}、成長期:{in_gro_name}、値の配列:{temp_value_list}") from e


def make_proactive_patterns(in_ptn_dict: dict[str, opc.OnePattern], is_large_minus: bool) -> dict[str, opc.OnePattern]:
    """積極鍛錬のパターンを作成

    Args:
        in_ptn_dict (dict[str, opc.OnePattern]): 1つの成長期のパターンの辞書
        is_large_minus (bool): 大練習マイナスがあるか

    Returns:
        dict[str, opc.OnePattern]: 新たに作成したパターンの辞書
    """
    # 新たに追加する経験値パターンを格納する辞書
    add_ex_pattern_dict: dict[str, opc.OnePattern] = {}

    for one_pattern_name, one_pattern in in_ptn_dict.items():

        # 加算処理フラグ(falseなら減算処理を実施)
        is_exec_add: bool = True

        if one_pattern_name == "large_train_ex":
            # 大練習の時
            add_one_pattern_name: str = "proactive_" + one_pattern_name
        elif "_ap_large_" in one_pattern_name and not one_pattern_name.startswith("mental_ap_"):
            # AP大練習（メンタリスト嫁除く）の時
            add_one_pattern_name = "proactive_" + one_pattern_name
        elif one_pattern_name == "small_train_ex":
            # 小練習の値を用いて、積極新球を作成
            add_one_pattern_name = "proactive_new_ball_large_train_ex"
        elif one_pattern_name == "large_minus_ex":
            # 大練習マイナスの時（減算処理を実行）
            add_one_pattern_name = "proactive_" + one_pattern_name
            is_exec_add = False
        else:
            # 上記のパターン以外は対応しない
            continue

        # 新しいパターンを作成し、加算ないし減算処理を実施
        add_ex_pattern_dict[add_one_pattern_name] = one_pattern.copy()
        if is_exec_add:
            add_ex_pattern_dict[add_one_pattern_name].add_proactive_ex_point(is_large_minus)
        else:
            add_ex_pattern_dict[add_one_pattern_name].sub_proactive_ex_point(is_large_minus)

    return add_ex_pattern_dict


def make_cautious_patterns(in_ptn_dict: dict[str, opc.OnePattern], is_large_minus: bool) -> dict[str, opc.OnePattern]:
    """慎重鍛錬のパターンを作成（大練習マイナスが対象）

    Args:
        in_ptn_dict (dict[str, opc.OnePattern]): 1つの成長期のパターンの辞書
        is_large_minus (bool): 大練習マイナスがあるか

    Returns:
        dict[str, opc.OnePattern]: 新たに作成したパターンの辞書
    """
    # 新たに追加する経験値パターンを格納する辞書
    add_ex_pattern_dict: dict[str, opc.OnePattern] = {}

    for one_pattern_name, one_pattern in in_ptn_dict.items():
        if one_pattern_name == "large_minus_ex":
            # 新しいパターンを作成し、除算処理を実施
            add_one_pattern_name: str = "cautious_" + one_pattern_name
            add_ex_pattern_dict[add_one_pattern_name] = one_pattern.copy()
            add_ex_pattern_dict[add_one_pattern_name].div_cautious_ex_point(is_large_minus)

    return add_ex_pattern_dict


def make_precise_patterns(in_ptn_dict: dict[str, opc.OnePattern]) -> dict[str, opc.OnePattern]:
    """精密鍛錬のパターンを作成（小練習すべてが対象）

    Args:
        in_ptn_dict (dict[str, opc.OnePattern]): 1つの成長期のパターンの辞書

    Returns:
        dict[str, opc.OnePattern]: 新たに作成したパターンの辞書
    """
    # 新たに追加する経験値パターンを格納する辞書
    add_ex_pattern_dict: dict[str, opc.OnePattern] = {}

    for one_pattern_name, one_pattern in in_ptn_dict.items():
        if "small_train" in one_pattern_name:
            # 新しいパターンを作成し、加算処理を実施
            add_one_pattern_name: str = "precise_" + one_pattern_name
            add_ex_pattern_dict[add_one_pattern_name] = one_pattern.copy()
            add_ex_pattern_dict[add_one_pattern_name].add_precise_ex_point()

    return add_ex_pattern_dict


def make_equilibrium_patterns(in_ptn_dict: dict[str, opc.OnePattern]) -> dict[str, opc.OnePattern]:
    """平衡鍛錬のパターンを作成（大練習と小練習すべてが対象）

    Args:
        in_ptn_dict (dict[str, opc.OnePattern]): 1つの成長期のパターンの辞書

    Returns:
        dict[str, opc.OnePattern]: 新たに作成したパターンの辞書
    """
    # 新たに追加する経験値パターンを格納する辞書
    add_ex_pattern_dict: dict[str, opc.OnePattern] = {}

    for one_pattern_name, one_pattern in in_ptn_dict.items():
        if ef.is_train_pattern(one_pattern_name):
            # 大練習か小練習なら、新しいパターンを作成し、加算処理を実施
            add_one_pattern_name: str = "equilibrium_" + one_pattern_name
            add_ex_pattern_dict[add_one_pattern_name] = one_pattern.copy()
            add_ex_pattern_dict[add_one_pattern_name].add_equilibrium_ex_point()

    return add_ex_pattern_dict


def add_training_patterns(in_ptn_dict: dict[str, opc.OnePattern], in_gro_name: str) -> None:
    """各鍛錬（積極・慎重・精密・平衡）のパターンを追加
    各鍛錬は、鍛錬適用前のパターンだけをもとに作成する。

    Args:
        in_ptn_dict (dict[str, opc.OnePattern]): 1つの成長期のパターンの辞書
        in_gro_name (str): 成長期名（エラーメッセージ用）

    Raises:
        ValueError: 大練習マイナスの値が空の場合
    """
    # 大練習マイナスが発生するかのフラグ
    is_large_minus: bool = has_large_minus(in_ptn_dict, in_gro_name)

    # すべての鍛錬のパターンを作成してから、まとめて追加する
    add_ex_pattern_dict: dict[str, opc.OnePattern] = {}
    add_ex_pattern_dict.update(make_proactive_patterns(in_ptn_dict, is_large_minus))
    add_ex_pattern_dict.update(make_cautious_patterns(in_ptn_dict, is_large_minus))
    add_ex_pattern_dict.update(make_precise_patterns(in_ptn_dict))
    add_ex_pattern_dict.update(make_equilibrium_patterns(in_ptn_dict))

    in_ptn_dict.update(add_ex_pattern_dict)


def apply_cast(in_ptn_dict: dict[str, opc.OnePattern]) -> None:
    """筋肉養成ギプスの適用（大練習と小練習全てに iv.CAST_MUL_FACTOR を乗算、端数は切り上げ）

    Args:
        in_ptn_dict (dict[str, opc.OnePattern]): 1つの成長期のパターンの辞書
    """
    for one_pattern_name, one_pattern in in_ptn_dict.items():
        if ef.is_train_pattern(one_pattern_name):
            one_pattern.mul_ex_point(iv.CAST_MUL_FACTOR, False)


def write_sheet(in_ws: Worksheet, in_all_ex_ptn_dict: dict[str, dict[str, opc.OnePattern]],
                in_output_type: int, in_sheet_name: str, in_output_name: str) -> None:
    """全ての成長期のパターンをシートに出力

    Args:
        in_ws (Worksheet): 出力先のシート
        in_all_ex_ptn_dict (dict[str, dict[str, opc.OnePattern]]): 成長期名をkeyとした、各成長期のパターンの辞書
        in_output_type (int): 0: 出現値、1: 期待値
        in_sheet_name (str): シートの本来の名前（乗っている補正）
        in_output_name (str): 出力内容の名前
    """
    excel_rows_val: int = 1

    for pattern_dict in in_all_ex_ptn_dict.values():
        for excel_column_val, column_name in enumerate(iv.OUTPUT_COLUMN_NAME_LIST, start=1):
            a_pattern_obj: opc.OnePattern = pattern_dict[column_name]
            if in_output_type == 1:
                # 出力対象が期待値の時
                write_val: str = str(a_pattern_obj.get_expected_value())
            else:
                # 出力対象が出現値の時
                write_val = format_values_to_str(a_pattern_obj.get_values_list())
            in_ws.cell(row=excel_rows_val, column=excel_column_val, value=write_val)
        excel_rows_val += 1

    # シートの本来の名前（乗っている補正）と、出力内容をそれぞれセルに出力
    in_ws.cell(row=excel_rows_val, column=1, value=in_sheet_name)
    in_ws.cell(row=excel_rows_val, column=2, value=in_output_name)


def format_values_to_str(in_values_list: list[int]) -> str:
    """出現値のリストをセルに出力する文字列に変換
    3つ以上連続する数がある場合は、その箇所は「～」で括る

    Args:
        in_values_list (list[int]): 出現値のリスト

    Returns:
        str: セルに出力する文字列（空配列なら "None"）
    """
    if len(in_values_list) == 0:
        # 空配列ならNoneを入れる
        return "None"
    if min(in_values_list) >= 0:
        # 配列が正の値か0のみなら、ソート処理を行う
        return ef.sort_and_omit_value_lists_to_str(in_values_list)

    # 配列に負の数があるなら、直接処理する（大練習マイナスはこの時点で連続しているため）
    if len(in_values_list) > 2:
        return str(max(in_values_list)) + "～" + str(min(in_values_list))
    if len(in_values_list) == 2:
        return str(max(in_values_list)) + "," + str(min(in_values_list))
    return str(in_values_list[0])
