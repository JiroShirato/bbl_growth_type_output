from openpyxl.worksheet.worksheet import Worksheet

import one_pattern_class as op_cls
import init_values as ini_value
import ext_functions as ext_f


def calc_growth_patterns(growth_ex_list: list[str], attributes: list[int]) -> dict[str, op_cls.OnePattern]:
    """1つの成長期について、シートの条件に応じた全ての経験値パターンを計算

    Args:
        growth_ex_list (list[str]): CSVの1行(成長期名と各練習の最小値・最大値)
        attributes (list[int]): シートの条件(小AP, YUR, 集中・イマイチ, ギプスの有無, 出力対象)

    Returns:
        dict[str, op_cls.OnePattern]: 経験値パターン名をkeyとした全ての経験値パターンの辞書

    Raises:
        ValueError: 数値に変換できない値がある、または大練習マイナスの値が空の場合
    """
    # シートの条件をそれぞれ取得
    mini_ap, yur, concentrate, gips, output_type = attributes

    # 成長期の名前
    growth_name: str = growth_ex_list[0]

    # 成長期における経験値パターンを格納する辞書(非AP小練習/非AP大練習/大練習マイナス/自主トレ参加 を最初に取得)
    patterns_dict: dict[str, op_cls.OnePattern] = create_base_patterns(growth_ex_list)

    # 1. 小APによる経験値加算処理
    if mini_ap == 1:
        add_ex_value_to_train_patterns(patterns_dict, ini_value.MINI_AP_ADD_VALUE)

    # 2. APによる経験値乗算処理
    add_ap_patterns(patterns_dict)

    # 3. YURによる経験値加算処理
    if yur == 1:
        add_ex_value_to_train_patterns(patterns_dict, ini_value.YUR_ADD_VALUE)

    # 4. メンタリスト能力持ちの嫁による経験値乗算処理
    add_mentalist_patterns(patterns_dict)

    # 5. 集中、及びイマイチやる気が出ない…による経験値加算、または除算処理
    if concentrate != 0:
        apply_concentrate(patterns_dict, concentrate)

    # 6. 各鍛錬による補正
    add_training_patterns(patterns_dict, growth_name)

    # 7. 筋肉養成ギプスの適用による経験値乗算補正
    if gips == 1:
        apply_gips(patterns_dict)

    # 8. 期待値算出(出力対象が期待値の時)
    if output_type == 1:
        for one_pattern in patterns_dict.values():
            one_pattern.calc_expected_ex_value()

    return patterns_dict


def create_base_patterns(growth_ex_list: list[str]) -> dict[str, op_cls.OnePattern]:
    """CSVの1行から、基本となる経験値パターン(非AP小練習、非AP大練習、大練習マイナス、自主トレ参加)を作成

    Args:
        growth_ex_list (list[str]): CSVの1行(成長期名と各練習の最小値・最大値)

    Returns:
        dict[str, op_cls.OnePattern]: 経験値パターン名をkeyとした経験値パターンの辞書
    """
    # 基本となる経験値パターンを格納する辞書
    patterns_dict: dict[str, op_cls.OnePattern] = {}

    # 非AP小練習の値の取得
    small_op_cls: op_cls.OnePattern = op_cls.OnePattern.from_strings(growth_ex_list[1], growth_ex_list[2])
    patterns_dict["small_train_ex"] = small_op_cls

    # 非AP大練習の値の取得
    large_op_cls: op_cls.OnePattern = op_cls.OnePattern.from_strings(growth_ex_list[3], growth_ex_list[4])
    patterns_dict["large_train_ex"] = large_op_cls

    # 大練習のマイナス値の取得
    minus_op_cls: op_cls.OnePattern = op_cls.OnePattern.from_strings(growth_ex_list[6], growth_ex_list[5])
    patterns_dict["large_minus_ex"] = minus_op_cls

    # 自主トレ参加の値の取得(非AP小練習の値を使用、非AP小練習の値がない場合はインスタンスだけ作る)
    small_op_cls_list: list[int] = small_op_cls.get_ex_values_list()
    if len(small_op_cls_list) == 0:
        patterns_dict["participate_independent_training_ex"] = op_cls.OnePattern()
    else:
        independent_op_cls: op_cls.OnePattern = \
            op_cls.OnePattern(min(small_op_cls_list) * 3, max(small_op_cls_list) * 3)
        patterns_dict["participate_independent_training_ex"] = independent_op_cls

    return patterns_dict


def add_ex_value_to_train_patterns(ptn_dict: dict[str, op_cls.OnePattern], add_ex_value: int) -> None:
    """大練習と小練習すべての経験値パターンに固定値を加算(小AP、YUR向け)

    Args:
        ptn_dict (dict[str, op_cls.OnePattern]): 1つの成長期の経験値パターンの辞書
        add_ex_value (int): 加算する値
    """
    for one_pattern_name, one_pattern in ptn_dict.items():
        if ext_f.is_train_pattern(one_pattern_name):
            one_pattern.add_ex_value(add_ex_value)


def add_ap_patterns(ptn_dict: dict[str, op_cls.OnePattern]) -> None:
    """APによる経験値倍増の経験値パターンを追加(倍率設定から、計算)

    Args:
        ptn_dict (dict[str, op_cls.OnePattern]): 1つの成長期の経験値パターンの辞書
    """
    # 新たに追加する経験値パターンを格納する辞書
    add_ptn_dict: dict[str, op_cls.OnePattern] = {}

    for ap_name, mul_factor in ini_value.AP_MULTIPLICATION_FACTOR_DICT.items():
        for one_pattern_name, one_pattern in ptn_dict.items():
            # 大練習と小練習すべてが対象
            if ext_f.is_train_pattern(one_pattern_name):
                # APの名前の設定
                add_one_pattern_name: str = ap_name + "_" + one_pattern_name
                # 新しい経験値パターンを作成し、APに設定された乗算値をかける(端数は切り捨て)
                add_ptn_dict[add_one_pattern_name] = one_pattern.copy()
                add_ptn_dict[add_one_pattern_name].mul_ex_value(mul_factor)

    ptn_dict.update(add_ptn_dict)


def add_mentalist_patterns(ptn_dict: dict[str, op_cls.OnePattern]) -> None:
    """メンタリスト能力持ちの嫁の経験値パターンを追加

    Args:
        ptn_dict (dict[str, op_cls.OnePattern]): 1つの成長期の経験値パターンの辞書
    """
    # 新たに追加する経験値パターンを格納する辞書
    add_ptn_dict: dict[str, op_cls.OnePattern] = {}

    for one_pattern_name, one_pattern in ptn_dict.items():
        if one_pattern_name == "large_train_ex" or one_pattern_name == "small_train_ex":
            # 大練習と小練習が対象なら、メンタリストを表す"mental_"の名前を追加
            add_one_pattern_name: str = "mental_" + one_pattern_name
        elif one_pattern_name.startswith("other_ap_"):
            # 1.5倍APを表す"other_ap_"なら、名前を精神APメンタリストを表す"mental_ap_"に変更
            add_one_pattern_name = one_pattern_name.replace("other_ap_", "mental_ap_")
        else:
            # 上記の経験値パターン以外は対応しない
            continue
        # 新しい経験値パターンを作成し、乗算値 ini_valuev.MENTALIST_WIFE_MUL_FACTOR をかける
        add_ptn_dict[add_one_pattern_name] = one_pattern.copy()
        add_ptn_dict[add_one_pattern_name].mul_ex_value(ini_value.MENTALIST_WIFE_MUL_FACTOR)

    ptn_dict.update(add_ptn_dict)


def apply_concentrate(ptn_dict: dict[str, op_cls.OnePattern], concentrate_type: int) -> None:
    """集中、及びイマイチやる気が出ない…による経験値加減処理

    Args:
        ptn_dict (dict[str, op_cls.OnePattern]): 1つの成長期の経験値パターンの辞書
        concentrate_type (int): 0: なし、1: 集中、2: イマイチ
    """
    for one_pattern_name, one_pattern in ptn_dict.items():
        # 大練習と小練習すべてが対象
        if ext_f.is_train_pattern(one_pattern_name):
            if concentrate_type == 1:
                # 集中の時
                one_pattern.add_concentrate_ex_values()
            elif concentrate_type == 2:
                # イマイチの時(2で割る)
                one_pattern.div_ex_values_with_ceil_and_floor(2)


def has_large_minus(ptn_dict: dict[str, op_cls.OnePattern], growth_name: str) -> bool:
    """大練習で経験値が下がるか(積極鍛錬と慎重鍛錬向け)

    Args:
        ptn_dict (dict[str, op_cls.OnePattern]): 1つの成長期の経験値パターンの辞書
        growth_name (str): 成長期名(エラーメッセージ用)

    Returns:
        bool: 大練習マイナスがあれば True

    Raises:
        ValueError: 大練習マイナスの値が空の場合
    """
    # 大練習マイナスの経験値パターンの値を格納したリスト
    large_minus_ex_values_list: list[int] = ptn_dict["large_minus_ex"].get_ex_values_list()
    try:
        return max(large_minus_ex_values_list) < 0
    except ValueError as e:
        raise ValueError(
            f"大練習マイナスの値が不正です: {e}、成長期:{growth_name}、値の配列:{large_minus_ex_values_list}") from e


def make_proactive_patterns(ptn_dict: dict[str, op_cls.OnePattern],
                            is_large_minus: bool) -> dict[str, op_cls.OnePattern]:
    """積極鍛錬の経験値パターンを作成

    Args:
        ptn_dict (dict[str, op_cls.OnePattern]): 1つの成長期の経験値パターンの辞書
        is_large_minus (bool): 大練習マイナスがあるか

    Returns:
        dict[str, op_cls.OnePattern]: 新たに作成した経験値パターンの辞書
    """
    # 新たに追加する経験値パターンを格納する辞書
    add_ptn_dict: dict[str, op_cls.OnePattern] = {}

    for one_pattern_name, one_pattern in ptn_dict.items():

        # 加算処理フラグ(falseなら減算処理を実施)
        is_exec_add: bool = True

        if one_pattern_name == "large_train_ex":
            # 大練習の時
            add_one_pattern_name: str = "proactive_" + one_pattern_name
        elif "_ap_large_" in one_pattern_name and not one_pattern_name.startswith("mental_ap_"):
            # AP大練習(メンタリスト嫁除く)の時
            add_one_pattern_name = "proactive_" + one_pattern_name
        elif one_pattern_name == "small_train_ex":
            # 小練習の値を用いて、積極新球を作成
            add_one_pattern_name = "proactive_new_ball_large_train_ex"
        elif one_pattern_name == "large_minus_ex":
            # 大練習マイナスの時(減算処理を実行)
            add_one_pattern_name = "proactive_" + one_pattern_name
            is_exec_add = False
        else:
            # 上記の経験値パターン以外は対応しない
            continue

        # 新しい経験値パターンを作成し、加算ないし減算処理を実施
        add_ptn_dict[add_one_pattern_name] = one_pattern.copy()
        if is_exec_add:
            add_ptn_dict[add_one_pattern_name].add_proactive_ex_values(is_large_minus)
        else:
            add_ptn_dict[add_one_pattern_name].sub_proactive_ex_values(is_large_minus)

    return add_ptn_dict


def make_cautious_patterns(ptn_dict: dict[str, op_cls.OnePattern],
                           is_large_minus: bool) -> dict[str, op_cls.OnePattern]:
    """慎重鍛錬の経験値パターンを作成(大練習マイナスが対象)

    Args:
        ptn_dict (dict[str, op_cls.OnePattern]): 1つの成長期の経験値パターンの辞書
        is_large_minus (bool): 大練習マイナスがあるか

    Returns:
        dict[str, op_cls.OnePattern]: 新たに作成した経験値パターンの辞書
    """
    # 新たに追加する経験値パターンを格納する辞書
    add_ptn_dict: dict[str, op_cls.OnePattern] = {}

    for one_pattern_name, one_pattern in ptn_dict.items():
        if one_pattern_name == "large_minus_ex":
            # 新しい経験値パターンを作成し、除算処理を実施
            add_one_pattern_name: str = "cautious_" + one_pattern_name
            add_ptn_dict[add_one_pattern_name] = one_pattern.copy()
            add_ptn_dict[add_one_pattern_name].div_cautious_ex_values(is_large_minus)

    return add_ptn_dict


def make_precise_patterns(ptn_dict: dict[str, op_cls.OnePattern]) -> dict[str, op_cls.OnePattern]:
    """精密鍛錬の経験値パターンを作成(小練習すべてが対象)

    Args:
        ptn_dict (dict[str, op_cls.OnePattern]): 1つの成長期の経験値パターンの辞書

    Returns:
        dict[str, op_cls.OnePattern]: 新たに作成した経験値パターンの辞書
    """
    # 新たに追加する経験値パターンを格納する辞書
    add_ptn_dict: dict[str, op_cls.OnePattern] = {}

    for one_pattern_name, one_pattern in ptn_dict.items():
        if "small_train" in one_pattern_name:
            # 小練習なら、新しい経験値パターンを作成し加算処理を実施
            add_one_pattern_name: str = "precise_" + one_pattern_name
            add_ptn_dict[add_one_pattern_name] = one_pattern.copy()
            add_ptn_dict[add_one_pattern_name].add_precise_ex_values()

    return add_ptn_dict


def make_equilibrium_patterns(ptn_dict: dict[str, op_cls.OnePattern]) -> dict[str, op_cls.OnePattern]:
    """平衡鍛錬の経験値パターンを作成(大練習と小練習すべてが対象)

    Args:
        ptn_dict (dict[str, op_cls.OnePattern]): 1つの成長期の経験値パターンの辞書

    Returns:
        dict[str, op_cls.OnePattern]: 新たに作成した経験値パターンの辞書
    """
    # 新たに追加する経験値パターンを格納する辞書
    add_ptn_dict: dict[str, op_cls.OnePattern] = {}

    for one_pattern_name, one_pattern in ptn_dict.items():
        if ext_f.is_train_pattern(one_pattern_name):
            # 大練習か小練習なら、新しい経験値パターンを作成し、加算処理を実施
            add_one_pattern_name: str = "equilibrium_" + one_pattern_name
            add_ptn_dict[add_one_pattern_name] = one_pattern.copy()
            add_ptn_dict[add_one_pattern_name].add_equilibrium_ex_values()

    return add_ptn_dict


def add_training_patterns(ptn_dict: dict[str, op_cls.OnePattern], growth_name: str) -> None:
    """各鍛錬(積極・慎重・精密・平衡)の経験値パターンを追加
    各鍛錬は、鍛錬適用前の経験値パターンだけを対象に作成する。

    Args:
        ptn_dict (dict[str, op_cls.OnePattern]): 1つの成長期の経験値パターンの辞書
        growth_name (str): 成長期名(エラーメッセージ用)

    Raises:
        ValueError: 大練習マイナスの値が空の場合
    """
    # 大練習マイナスが発生するかのフラグ
    is_large_minus: bool = has_large_minus(ptn_dict, growth_name)

    # すべての鍛錬の経験値パターンを作成してから、まとめて追加する
    add_ptn_dict: dict[str, op_cls.OnePattern] = {}
    add_ptn_dict.update(make_proactive_patterns(ptn_dict, is_large_minus))
    add_ptn_dict.update(make_cautious_patterns(ptn_dict, is_large_minus))
    add_ptn_dict.update(make_precise_patterns(ptn_dict))
    add_ptn_dict.update(make_equilibrium_patterns(ptn_dict))

    ptn_dict.update(add_ptn_dict)


def apply_gips(ptn_dict: dict[str, op_cls.OnePattern]) -> None:
    """筋肉養成ギプスの適用(大練習と小練習全てに ini_value.GIPS_MUL_FACTOR を乗算、端数は切り上げ)

    Args:
        ptn_dict (dict[str, op_cls.OnePattern]): 1つの成長期の経験値パターンの辞書
    """
    for one_pattern_name, one_pattern in ptn_dict.items():
        if ext_f.is_train_pattern(one_pattern_name):
            one_pattern.mul_ex_value(ini_value.GIPS_MUL_FACTOR, False)


def write_sheet(ws: Worksheet, all_ptn_dict: dict[str, dict[str, op_cls.OnePattern]],
                output_type: int, sheet_name: str, output_name: str) -> None:
    """全ての成長期の経験値パターンをシートに出力

    Args:
        ws (Worksheet): 出力先のシート
        all_ptn_dict (dict[str, dict[str, op_cls.OnePattern]]):
            成長期名をkeyとした、各成長期の経験値パターンの辞書
        output_type (int): 0: 出現値、1: 期待値
        sheet_name (str): シートの本来の名前(乗っている補正)
        output_name (str): 出力内容の名前
    """
    excel_row: int = 1

    for pattern_dict in all_ptn_dict.values():
        for excel_column, column_name in enumerate(ini_value.OUTPUT_COLUMN_NAME_LIST, start=1):
            a_pattern_obj: op_cls.OnePattern = pattern_dict[column_name]
            if output_type == 1:
                # 出力対象が期待値の時
                write_ex_value: str = str(a_pattern_obj.get_expected_ex_value())
            else:
                # 出力対象が出現値の時
                write_ex_value = format_ex_values_to_str(a_pattern_obj.get_ex_values_list())
            ws.cell(row=excel_row, column=excel_column, value=write_ex_value)
        excel_row += 1

    # シートの本来の名前(乗っている補正)と、出力内容をそれぞれセルに出力
    ws.cell(row=excel_row, column=1, value=sheet_name)
    ws.cell(row=excel_row, column=2, value=output_name)


def format_ex_values_to_str(ex_values_list: list[int]) -> str:
    """出現値のリストをセルに出力する文字列に変換
    3つ以上連続する数がある場合は、その箇所は「～」で括る

    Args:
        ex_values_list (list[int]): 出現値のリスト

    Returns:
        str: セルに出力する文字列(空配列なら "None")
    """
    if len(ex_values_list) == 0:
        # 空配列ならNoneを入れる
        return "None"
    if min(ex_values_list) >= 0:
        # 配列が正の値か0のみなら、ソート処理を行う
        return ext_f.sort_and_omit_ex_values_list_to_str(ex_values_list)

    # 配列に負の数があるなら、直接処理する(大練習マイナスはこの時点で連続しているため)
    if len(ex_values_list) > 2:
        return str(max(ex_values_list)) + "～" + str(min(ex_values_list))
    if len(ex_values_list) == 2:
        return str(max(ex_values_list)) + "," + str(min(ex_values_list))
    return str(ex_values_list[0])
