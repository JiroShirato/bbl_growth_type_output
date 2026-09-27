import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

# プロジェクトのルートを import のパスに追加（テストの実行場所に依存しないようにする）
PROJECT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_DIR))

import main  # noqa: E402

# 入力CSV
INPUT_CSV_PATH = PROJECT_DIR / "import.csv"

# 比較の基準となる出力Excel（import.csv から出力した、正しいと確認済みの結果）
EXPECTED_EXCEL_PATH = Path(__file__).resolve().parent / "expected_output.xlsx"


def read_sheet_values(worksheet) -> list[list]:
    """シートの全セルの値を、末尾の空の行と列を除いた2次元リストで返す

    Excelで保存し直したファイルには書式だけの空セルが含まれることがあるため、
    値の比較に影響しないよう取り除く。

    Args:
        worksheet: openpyxl のワークシート

    Returns:
        list[list]: 各行のセルの値のリスト
    """
    rows: list[list] = [list(row) for row in worksheet.iter_rows(values_only=True)]

    # 各行の末尾の空セルを除く
    for row in rows:
        while row and row[-1] is None:
            row.pop()

    # 末尾の空行を除く
    while rows and not rows[-1]:
        rows.pop()

    return rows


class TestRegression(unittest.TestCase):
    """import.csv の出力結果が、基準の出力Excelと一致するかを確認する回帰テスト"""

    def test_output_matches_expected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "output.xlsx"
            main.main(input_file=str(INPUT_CSV_PATH), output_file=str(output_path))

            expected_wb = load_workbook(EXPECTED_EXCEL_PATH, read_only=True)
            actual_wb = load_workbook(output_path, read_only=True)
            try:
                # シート名と並び順
                self.assertEqual(expected_wb.sheetnames, actual_wb.sheetnames)

                # 各シートの全セル（値のある行数・列数の違いも差分として検出する）
                for sheet_name in expected_wb.sheetnames:
                    with self.subTest(sheet=sheet_name):
                        expected_rows = read_sheet_values(expected_wb[sheet_name])
                        actual_rows = read_sheet_values(actual_wb[sheet_name])
                        self.assertEqual(expected_rows, actual_rows)
            finally:
                expected_wb.close()
                actual_wb.close()


if __name__ == "__main__":
    unittest.main()
